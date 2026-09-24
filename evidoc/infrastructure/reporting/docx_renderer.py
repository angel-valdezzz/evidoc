"""Editable Word rendering of the same evidence model."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor
from reportlab.lib.utils import ImageReader

from evidoc.application.report_filename import safe_name
from evidoc.application.report_renderer import ReportRenderer
from evidoc.domain.artifact_type import ArtifactType
from evidoc.domain.run import Run
from evidoc.infrastructure.reporting.helpers import (
    fitted_size,
    image_bytes,
    report_date,
    status_color,
    summary_rows,
)


class DocxReportRenderer(ReportRenderer):
    format_name = "docx"

    def render_single(self, source_dir: Path, output_dir: Path, result: Run) -> Path:
        output = output_dir / f"{safe_name(result.test_case.name)}.docx"
        self._build(output, source_dir, [result])
        return output

    def render_run(
        self,
        source_dir: Path,
        output_dir: Path,
        results: list[Run],
        sources: dict[str, Path] | None = None,
    ) -> Path:
        run_ids = {result.run_id for result in results}
        run_id = results[0].run_id if len(run_ids) == 1 else "combined"
        output = output_dir / f"run-{run_id}.docx"
        self._build(output, source_dir, results, sources)
        return output

    def _build(
        self,
        output: Path,
        source_dir: Path,
        results: list[Run],
        sources: dict[str, Path] | None = None,
    ) -> None:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.8)
        header = section.header.paragraphs[0]
        header.paragraph_format.tab_stops.add_tab_stop(Inches(6.7), WD_TAB_ALIGNMENT.RIGHT)
        header.text = (
            f"{results[0].test_case.brand or 'EviDoc'}\t{results[0].test_case.project or ''}"
        )
        header.runs[0].bold = True
        header.runs[0].font.color.rgb = RGBColor(36, 42, 53)
        footer = section.footer.paragraphs[0]
        footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.7), WD_TAB_ALIGNMENT.RIGHT)
        footer.text = f"Ambiente: {results[0].test_case.environment or '—'}\tPágina "
        page_field = OxmlElement("w:fldSimple")
        page_field.set(qn("w:instr"), "PAGE")
        footer._p.append(page_field)
        for index, result in enumerate(results):
            date = document.add_paragraph(f"Fecha: {report_date(result)}")
            if index:
                date.paragraph_format.page_break_before = True
            date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            document.add_heading("Reporte de Ejecución Automatizada", 0)
            document.add_heading("Resumen De Ejecución", 1)
            table = document.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for label, value in summary_rows(result):
                row = table.add_row()
                row.cells[0].text = label
                row.cells[1].text = value
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "C62828" if label == "Defecto" else "2F50C5")
                row.cells[0]._tc.get_or_add_tcPr().append(shading)
                for run in row.cells[0].paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            artifacts = {artifact.id: artifact for artifact in result.artifacts}
            image_count = 0
            for step_index, step in enumerate(result.steps):
                break_before = image_count == 2 and any(
                    artifacts[identifier].type == ArtifactType.IMAGE
                    for identifier in step.artifact_ids
                )
                if break_before:
                    image_count = 0
                heading = document.add_heading(level=2)
                heading.paragraph_format.keep_with_next = True
                heading.paragraph_format.page_break_before = step_index == 0 or break_before
                marker = heading.add_run("◆ ")
                marker.font.color.rgb = RGBColor.from_string(status_color(step.status)[1:])
                heading.add_run(step.title)
                marker = heading.add_run(" ◆")
                marker.font.color.rgb = RGBColor.from_string(status_color(step.status)[1:])
                for log in step.logs:
                    document.add_paragraph(log.message)
                for identifier in step.artifact_ids:
                    artifact = artifacts[identifier]
                    if artifact.type != ArtifactType.IMAGE:
                        document.add_paragraph(
                            "Adjunto: " + (artifact.title or artifact.path or identifier)
                        )
                        continue
                    if image_count == 2:
                        next_picture_break = True
                        image_count = 0
                    else:
                        next_picture_break = False
                    image_count += 1
                    if artifact.title and artifact.title != step.title:
                        caption = document.add_paragraph(artifact.title)
                        caption.paragraph_format.keep_with_next = True
                        caption.paragraph_format.page_break_before = next_picture_break
                        next_picture_break = False
                    data = image_bytes(
                        (sources or {}).get(result.test_id, source_dir), result, artifact
                    )
                    if data:
                        paragraph = document.add_paragraph()
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        paragraph.paragraph_format.page_break_before = next_picture_break
                        paragraph.paragraph_format.keep_with_next = bool(artifact.description)
                        width, height = ImageReader(BytesIO(data)).getSize()
                        max_height = 8.5 if artifact.orientation == "vertical" else 5.5
                        display_width, display_height = fitted_size(width, height, 6.2, max_height)
                        paragraph.add_run().add_picture(
                            BytesIO(data),
                            width=Inches(display_width),
                            height=Inches(display_height),
                        )
                    else:
                        document.add_paragraph("Imagen no disponible")
                    if artifact.description:
                        document.add_paragraph(artifact.description)
        document.save(str(output))
