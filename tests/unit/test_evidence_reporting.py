from __future__ import annotations

import json
from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.shared import RGBColor
from evidoc import build, merge
from evidoc.api import EvidocAPI
from evidoc.infrastructure.filesystem.filesystem_result_repository import FilesystemResultRepository
from evidoc.infrastructure.reporting.helpers import fitted_size
from evidoc.interfaces.cli.main import app
from evidoc.listener import Listener
from evidoc.schema_paths import result_schema_path
from PIL import Image
from robot.run import run
from typer.testing import CliRunner


def png() -> bytes:
    stream = BytesIO()
    Image.new("RGB", (80, 50), "blue").save(stream, format="PNG")
    return stream.getvalue()


def write_worker(root: Path, number: int) -> None:
    api = EvidocAPI(root_dir=root)
    api.start_test(f"Process {number}")
    api.capture_image(png(), title=f"Evidence {number}")
    api.end_test("PASS", 0.1)


@pytest.mark.parametrize("storage", ["file", "base64"])
def test_storage_and_both_renderers(tmp_path: Path, storage: str) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source, storage=storage, application="Portal", requirement="REQ-9")
    api.start_test("Inicio de sesión")
    for kind in ("page", "element", "desktop"):
        assert api.capture_image(png(), title=kind, capture=kind, orientation="horizontal")
    api.end_test("PASS", 1.5)
    results = FilesystemResultRepository(result_schema_path()).load_test_results(source)
    assert len(results) == 1
    assert [step.title for step in results[0].steps] == ["page", "element", "desktop"]
    result = results[0]
    assert all(bool(a.data) == (storage == "base64") for a in result.artifacts)
    assert len(list(source.rglob("*.png"))) == (0 if storage == "base64" else 3)
    outputs = build(source, tmp_path / "reports", ["pdf", "docx"])
    assert {output.name for output in outputs} == {"Inicio_de_sesi_n.pdf", "Inicio_de_sesi_n.docx"}
    assert {output.suffix for output in outputs} == {".pdf", ".docx"}
    assert all(output.stat().st_size > 1000 for output in outputs)
    document = Document(str(next(p for p in outputs if p.suffix == ".docx")))
    assert "Reporte de Ejecución Automatizada" in [p.text for p in document.paragraphs]
    assert "Portal" in [cell.text for row in document.tables[0].rows for cell in row.cells]
    assert len(document.inline_shapes) == 3
    assert "INFO" not in [paragraph.text for paragraph in document.paragraphs]
    assert [paragraph.text for paragraph in document.paragraphs if "◆" in paragraph.text] == [
        "◆ page ◆",
        "◆ element ◆",
        "◆ desktop ◆",
    ]
    assert [row.cells[0].text for row in document.tables[0].rows][-1] == "Defecto"
    assert document.tables[0].rows[-1].cells[1].text == ""
    assert 'w:fill="C62828"' in document.tables[0].rows[-1].cells[0]._tc.xml
    from pypdf import PdfReader

    pdf_text = "\n".join(
        page.extract_text()
        for page in PdfReader(next(p for p in outputs if p.suffix == ".pdf")).pages
    )
    assert "Resumen De Ejecución" in pdf_text
    assert "INFO" not in pdf_text
    assert pdf_text.index("page") < pdf_text.index("element") < pdf_text.index("desktop")


def test_report_markers_follow_step_status(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    api.start_test("Status colors")
    api.log_step("Needs attention", "WARN")
    api.end_test("PASS", 0.2)
    outputs = build(source, tmp_path / "reports", ["pdf", "docx"])
    document = Document(str(next(path for path in outputs if path.suffix == ".docx")))
    heading = next(
        paragraph for paragraph in document.paragraphs if "Needs attention" in paragraph.text
    )
    assert heading.text == "◆ Needs attention ◆"
    assert [run.font.color.rgb for run in (heading.runs[0], heading.runs[-1])] == [
        RGBColor(237, 108, 2),
        RGBColor(237, 108, 2),
    ]
    from pypdf import PdfReader

    pdf = PdfReader(next(path for path in outputs if path.suffix == ".pdf"))
    pdf_text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Needs attention" in pdf_text
    assert "WARN" not in pdf_text
    contents = [page.get_contents() for page in pdf.pages]
    assert all(content is not None for content in contents)
    assert b".929412 .423529 .007843 rg" in b"\n".join(
        content.get_data() for content in contents if content is not None
    )


def test_same_case_name_cannot_overwrite_a_report(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    for _ in range(2):
        api.start_test("Same case")
        api.end_test("PASS", 0.1)
    with pytest.raises(ValueError, match="Case names must be unique"):
        build(source, tmp_path / "reports", ["pdf"])
    assert not (tmp_path / "reports").exists()


def test_defects_apply_only_to_named_cases_during_build(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    for name in ("TC-01", "TC-02"):
        api.start_test(name, full_name=f"Suite.{name}")
        api.end_test("PASS", 0.1)
    with pytest.raises(ValueError, match="exactly one selected test"):
        build(source, tmp_path / "rejected", defects=["BUG-100"])
    assert not (tmp_path / "rejected").exists()
    invalid_cli = CliRunner().invoke(
        app, ["build", "--input-dir", str(source), "--defect", "BUG-100"]
    )
    assert invalid_cli.exit_code != 0
    assert "exactly one selected test" in invalid_cli.output
    outputs = build(
        source,
        tmp_path / "reports",
        formats=["docx"],
        defects=["TC-01=BUG-100", "Suite.TC-01=BUG-101"],
    )
    defect_rows = {
        path.stem: Document(str(path)).tables[0].rows[-1].cells[1].text for path in outputs
    }
    assert defect_rows == {"TC-01": "BUG-100, BUG-101", "TC-02": ""}
    assert all(
        result.test_case.defect is None
        for result in FilesystemResultRepository(result_schema_path()).load_test_results(source)
    )


def test_single_selected_case_accepts_unqualified_defects(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    api.start_test("Passing case")
    api.end_test("PASS", 0.1)
    api.start_test("Failing case")
    api.end_test("FAIL", 0.1)
    outputs = build(
        source,
        tmp_path / "reports",
        formats=["docx"],
        exclude_status="FAIL",
        defects=["BUG-200", "BUG-201"],
    )
    assert len(outputs) == 1
    assert Document(str(outputs[0])).tables[0].rows[-1].cells[1].text == "BUG-200, BUG-201"


def test_defect_rejects_unknown_and_ambiguous_names(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    for suite in ("A", "B"):
        api.start_test("Same name", full_name=f"{suite}.Same name")
        api.end_test("PASS", 0.1)
    for defect, expected in (("Unknown=BUG-1", "matched 0"), ("Same name=BUG-1", "matched 2")):
        with pytest.raises(ValueError, match=expected):
            build(source, tmp_path / "reports", defects=[defect])
    assert not (tmp_path / "reports").exists()


def test_cli_repeated_defect_options_are_scoped_to_case(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    for name in ("TC036", "TC037"):
        api.start_test(name)
        api.end_test("PASS", 0.1)
    result = CliRunner().invoke(
        app,
        [
            "build",
            "--input-dir",
            str(source),
            "--output-dir",
            str(tmp_path / "reports"),
            "--formats",
            "docx",
            "--defect",
            "TC036=BUG-123",
            "--defect",
            "TC037=BUG-456",
        ],
    )
    assert result.exit_code == 0, result.stdout
    assert (
        Document(str(tmp_path / "reports" / "TC036.docx")).tables[0].rows[-1].cells[1].text
        == "BUG-123"
    )
    assert (
        Document(str(tmp_path / "reports" / "TC037.docx")).tables[0].rows[-1].cells[1].text
        == "BUG-456"
    )


def test_cli_and_python_build_match(tmp_path: Path) -> None:
    source = tmp_path / "metadata"
    api = EvidocAPI(root_dir=source)
    api.start_test("Caso")
    api.capture_image(png(), title="Pantalla")
    api.end_test("PASS", 0.2)
    expected = build(source, tmp_path / "python", ["pdf", "docx"])
    result = CliRunner().invoke(
        app,
        [
            "build",
            "--input-dir",
            str(source),
            "--output-dir",
            str(tmp_path / "cli"),
            "--formats",
            "pdf,docx",
        ],
    )
    assert result.exit_code == 0, result.stdout
    assert {p.name for p in expected} == {
        p.name for p in (tmp_path / "cli").iterdir() if p.suffix in {".pdf", ".docx"}
    }
    assert json.loads((tmp_path / "cli" / "upload-manifest.json").read_text()) == {
        "tests": [
            {
                "name": "Caso",
                "files": [str(path.resolve()) for path in (tmp_path / "cli").glob("*.pdf")]
                + [str(path.resolve()) for path in (tmp_path / "cli").glob("*.docx")],
            }
        ]
    }


def test_merge_selects_last_complete_attempt_and_filters_afterward(tmp_path: Path) -> None:
    original = tmp_path / "run" / "metadata"
    rerun = tmp_path / "rerun" / "metadata"
    first = EvidocAPI(root_dir=original)
    for name, status in (("TC-01", "PASS"), ("TC-02", "FAIL"), ("TC-03", "FAIL")):
        first.start_test(name, full_name=f"Suite.{name}")
        first.capture_image(png(), title=f"Run {name}")
        first.end_test(status, 0.1)
    replacement = EvidocAPI(root_dir=rerun)
    for name, status in (("TC-02", "PASS"), ("TC-03", "FAIL")):
        replacement.start_test(name, full_name=f"Suite.{name}")
        replacement.capture_image(png(), title=f"Rerun {name}")
        replacement.end_test(status, 0.1)

    index = merge([original, rerun], tmp_path / "final" / "metadata")
    assert index.name == "merged-results.json"
    assert not list(index.parent.rglob("*.png"))
    loaded = FilesystemResultRepository(result_schema_path()).load_test_results(index.parent)
    assert len(loaded) == 3
    assert [step.title for step in loaded[1].steps] == ["Rerun TC-02"]
    assert loaded[2].test_case.status.value == "FAIL"
    outputs = build(
        index.parent,
        tmp_path / "final" / "reports",
        ["pdf", "docx"],
        exclude_status=["FAIL", "SKIP"],
    )
    assert {path.name for path in outputs} == {"TC-01.pdf", "TC-01.docx", "TC-02.pdf", "TC-02.docx"}
    manifest = json.loads((tmp_path / "final" / "reports" / "upload-manifest.json").read_text())
    assert [case["name"] for case in manifest["tests"]] == ["TC-01", "TC-02"]
    assert all(Path(path).is_file() for case in manifest["tests"] for path in case["files"])


def test_merge_distinguishes_same_case_name_in_different_suites(tmp_path: Path) -> None:
    first = tmp_path / "run"
    second = tmp_path / "rerun"
    for directory, suite, status in (
        (first, "Suite A", "PASS"),
        (first, "Suite B", "FAIL"),
        (second, "Suite B", "PASS"),
    ):
        api = EvidocAPI(root_dir=directory)
        api.start_test("Same name", full_name=f"{suite}.Same name")
        api.end_test(status, 0.1)
    index = merge([first, second], tmp_path / "final")
    results = FilesystemResultRepository(result_schema_path()).load_test_results(index.parent)
    assert [(case.test_case.full_name, case.test_case.status.value) for case in results] == [
        ("Suite A.Same name", "PASS"),
        ("Suite B.Same name", "PASS"),
    ]


def test_attach_file_uses_original_path_without_copy(tmp_path: Path) -> None:
    metadata = tmp_path / "metadata"
    source = tmp_path / "downloads" / "caratula.pdf"
    source.parent.mkdir()
    source.write_bytes(b"example")
    api = EvidocAPI(root_dir=metadata, storage="base64")
    api.start_test("Policy")
    assert api.attach_file(source, "Carátula")
    api.end_test("PASS", 1)
    assert not list(metadata.rglob("*.pdf"))
    assert not list(metadata.rglob("*.png"))
    reports = build(metadata, tmp_path / "reports", ["pdf"])
    manifest = json.loads((tmp_path / "reports" / "upload-manifest.json").read_text())
    assert manifest["tests"] == [
        {"name": "Policy", "files": [str(reports[0].resolve()), str(source.resolve())]}
    ]
    source.unlink()
    with pytest.raises(FileNotFoundError, match="Attached file is missing"):
        build(metadata, tmp_path / "again", ["pdf"])
    assert not (tmp_path / "again").exists()


def test_cli_merge_and_single_status_exclusion(tmp_path: Path) -> None:
    source = tmp_path / "run"
    api = EvidocAPI(root_dir=source)
    for name, status in (("Passed", "PASS"), ("Skipped", "SKIP")):
        api.start_test(name, full_name=f"Suite.{name}")
        api.end_test(status, 0.1)
    result = CliRunner().invoke(
        app, ["merge", "--input-dir", str(source), "--output-dir", str(tmp_path / "merged")]
    )
    assert result.exit_code == 0, result.stdout
    result = CliRunner().invoke(
        app,
        [
            "build",
            "--input-dir",
            str(tmp_path / "merged"),
            "--output-dir",
            str(tmp_path / "reports"),
            "--exclude-status",
            "SKIP",
        ],
    )
    assert result.exit_code == 0, result.stdout
    assert {path.name for path in (tmp_path / "reports").iterdir()} == {
        "Passed.pdf",
        "upload-manifest.json",
    }
    assert {
        path.name for path in build(source, tmp_path / "plain", ["pdf"], exclude_status="SKIP")
    } == {"Passed.pdf"}


def test_robot_run_listener_and_capture_adapters(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class Element:
        screenshot_as_png = png()

    class Selenium:
        driver = Element()

        def find_element(self, locator: str) -> Element:
            assert locator == "//div[@id='PanelTitular']"
            return Element()

    monkeypatch.setattr("evidoc.robot.BuiltIn.get_library_instance", lambda self, name: Selenium())
    monkeypatch.setattr("evidoc.robot.desktop_bytes", png)
    download = tmp_path / "quote.pdf"
    download.write_bytes(b"downloaded")
    suite = tmp_path / "suite.robot"
    suite.write_text(
        f"""*** Settings ***
Library    evidoc.robot
*** Test Cases ***
Capture all
    Capture Page Evidence    Credenciales ingresadas    INFO    description=Folio 7
    Capture Element Evidence    //div[@id='PanelTitular']    Panel Titular    INFO    orientation=horizontal    include_page=True    description=Solo panel
    Capture Desktop Evidence    Evidencia completa    INFO    description=Todo el escritorio
    Attach File    {download}    description=Cotización
""",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    assert (
        run(
            str(suite), listener="evidoc.listener", outputdir=str(output), log="NONE", report="NONE"
        )
        == 0
    )
    results = FilesystemResultRepository(result_schema_path()).load_test_results(
        output / "evidoc" / "metadata"
    )
    assert len(results) == 1
    assert [a.capture for a in results[0].artifacts] == ["page", "page", "element", "desktop", None]
    assert [a.description for a in results[0].artifacts[:4]] == [
        "Folio 7",
        None,
        "Solo panel",
        "Todo el escritorio",
    ]
    assert results[0].artifacts[2].orientation == "horizontal"
    assert results[0].steps[1].title == "Contexto: Panel Titular"
    assert results[0].test_case.full_name is not None
    assert results[0].test_case.full_name.endswith("Capture all")
    assert results[0].artifacts[4].external
    assert results[0].artifacts[4].path == str(download.resolve())
    assert not list((output / "evidoc" / "metadata").rglob("*.pdf"))


def test_resource_imports_evidoc_keywords_with_listener(tmp_path: Path) -> None:
    resource = tmp_path / "evidencia.resource"
    resource.write_text(
        "*** Settings ***\nLibrary    evidoc.robot\n*** Keywords ***\n"
        "Registrar evidencia\n    Log Step    Inicio del flujo    INFO\n"
        "    Log Info    Dato de negocio\n",
        encoding="utf-8",
    )
    suite = tmp_path / "caso.robot"
    suite.write_text(
        "*** Settings ***\nResource    evidencia.resource\n*** Test Cases ***\n"
        "Caso con resource\n    Registrar evidencia\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    assert (
        run(
            str(suite), listener="evidoc.listener", outputdir=str(output), log="NONE", report="NONE"
        )
        == 0
    )
    results = FilesystemResultRepository(result_schema_path()).load_test_results(
        output / "evidoc" / "metadata"
    )
    assert len(results) == 1
    assert results[0].steps[0].title == "Inicio del flujo"


def test_small_element_uses_available_report_space() -> None:
    assert fitted_size(80, 50, 480, 340) == (480, 300)
    assert fitted_size(800, 50, 480, 340) == (480, 30)


def test_no_more_than_two_images_per_page_and_descriptions_follow_images(
    tmp_path: Path,
) -> None:
    metadata = tmp_path / "metadata"
    api = EvidocAPI(root_dir=metadata)
    api.start_test("Five captures")
    for number in range(5):
        api.capture_image(png(), title=f"Image {number}", description=f"Note {number}")
    api.end_test("PASS", 1)
    reports = build(metadata, tmp_path / "reports", ["pdf", "docx"])
    from pypdf import PdfReader

    pdf = PdfReader(next(report for report in reports if report.suffix == ".pdf"))
    texts = [page.extract_text() for page in pdf.pages]
    assert sum(text.count("♦ Image") for text in texts) == 5
    assert all(text.count("♦ Image") <= 2 for text in texts)
    assert all(
        any(f"♦ Image {number} ♦" in text and f"Note {number}" in text for text in texts)
        for number in range(5)
    )
    docx = Document(str(next(report for report in reports if report.suffix == ".docx")))
    assert [
        paragraph.text
        for paragraph in docx.paragraphs
        if paragraph.paragraph_format.page_break_before
    ] == ["◆ Image 0 ◆", "◆ Image 2 ◆", "◆ Image 4 ◆"]
    for number in range(5):
        note = next(i for i, p in enumerate(docx.paragraphs) if p.text == f"Note {number}")
        assert "<wp:inline" in docx.paragraphs[note - 1]._p.xml


def test_toml_config_applies_to_listener_and_build(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    (tmp_path / "evidoc.toml").write_text(
        'aplicacion = "Salud"\nproyecto = "Espartaco"\nenvironment = "QA"\nbrand = "AXA"\n'
        'formats = ["pdf", "docx"]\nstorage = "base64"\nmetadata_dir = "metadata"\n'
        'output_dir = "reports"\n',
        encoding="utf-8",
    )
    monkeypatch.setattr("evidoc.robot.desktop_bytes", png)
    suite = tmp_path / "test.robot"
    suite.write_text(
        "*** Settings ***\nLibrary    evidoc.robot\n*** Test Cases ***\nCaso real\n"
        "    Capture Desktop Evidence    Evidencia\n",
        encoding="utf-8",
    )
    assert (
        run(
            str(suite),
            listener=Listener(),
            outputdir=str(tmp_path / "robot"),
            log="NONE",
            report="NONE",
        )
        == 0
    )
    results = FilesystemResultRepository(result_schema_path()).load_test_results(
        tmp_path / "metadata"
    )
    assert len(results) == 1
    assert results[0].test_case.application == "Salud"
    assert results[0].test_case.project == "Espartaco"
    assert results[0].test_case.defect is None
    assert results[0].artifacts[0].data
    outputs = build()
    assert {output.suffix for output in outputs} == {".pdf", ".docx"}
    assert all(output.parent.resolve() == tmp_path / "reports" for output in outputs)
    cli = CliRunner().invoke(app, ["build", "--output-dir", str(tmp_path / "cli-reports")])
    assert cli.exit_code == 0, cli.stdout
    assert {output.name for output in outputs} == {
        output.name
        for output in (tmp_path / "cli-reports").iterdir()
        if output.suffix in {".pdf", ".docx"}
    }


def test_robot_listener_storage_configuration(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr("evidoc.robot.desktop_bytes", png)
    suite = tmp_path / "desktop.robot"
    suite.write_text(
        "*** Settings ***\nLibrary    evidoc.robot\n*** Test Cases ***\nCase\n    Capture Desktop Evidence    Escritorio\n",
        encoding="utf-8",
    )
    metadata = tmp_path / "shared" / "metadata"
    assert (
        run(
            str(suite),
            listener=Listener(root_dir=str(metadata), storage="base64"),
            outputdir=str(tmp_path / "robot"),
            log="NONE",
            report="NONE",
        )
        == 0
    )
    results = FilesystemResultRepository(result_schema_path()).load_test_results(metadata)
    assert len(results) == 1
    assert results[0].artifacts[0].data
    assert not list(metadata.rglob("*.png"))


def test_missing_capture_adapter_warns_without_failing_robot(tmp_path: Path) -> None:
    suite = tmp_path / "no_browser.robot"
    suite.write_text(
        "*** Settings ***\nLibrary    evidoc.robot\n*** Test Cases ***\nCase\n    Capture Page Evidence    Sin navegador\n",
        encoding="utf-8",
    )
    output = tmp_path / "output"
    assert (
        run(
            str(suite), listener="evidoc.listener", outputdir=str(output), log="NONE", report="NONE"
        )
        == 0
    )
    results = FilesystemResultRepository(result_schema_path()).load_test_results(
        output / "evidoc" / "metadata"
    )
    assert len(results) == 1
    assert not results[0].artifacts


def test_concurrent_writers_share_run_without_overwriting(tmp_path: Path) -> None:
    root = tmp_path / "metadata"

    def worker(i: int) -> None:
        api = EvidocAPI(root_dir=root)
        api.start_test(f"Case {i}")
        api.capture_image(png(), title=f"Evidence {i}")
        api.end_test("PASS", 0.1)

    with ThreadPoolExecutor(max_workers=8) as pool:
        list(pool.map(worker, range(24)))
    results = FilesystemResultRepository(result_schema_path()).load_test_results(root)
    assert len(results) == 24
    assert len({r.run_id for r in results}) == 1
    assert len({r.test_id for r in results}) == 24


def test_separate_processes_share_metadata_directory(tmp_path: Path) -> None:
    root = tmp_path / "metadata"
    with ProcessPoolExecutor(max_workers=4) as pool:
        list(pool.map(write_worker, [root] * 12, range(12)))
    results = FilesystemResultRepository(result_schema_path()).load_test_results(root)
    assert len(results) == 12
    assert len({result.run_id for result in results}) == 1
